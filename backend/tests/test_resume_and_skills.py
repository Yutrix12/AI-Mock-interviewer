import io

import pytest
from docx import Document

from backend import skills
from backend.profile import build_profile, merge_skills, seniority_for
from backend.resume import ResumeError, extract_text

from .conftest import SAMPLE_RESUME, FakeLLM, resume_draft
from backend.schemas import ResumeDraft


def _minimal_pdf(text: str) -> bytes:
    """A tiny one-page PDF with a single line of text, built by hand."""
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO(b"%PDF-1.4\n")
    offsets = []
    for i, obj in enumerate(objects, 1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n" % i + obj + b"\nendobj\n")
    xref = out.tell()
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1))
    for off in offsets:
        out.write(b"%010d 00000 n \n" % off)
    out.write(b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF" % (len(objects) + 1, xref))
    return out.getvalue()


def test_extracts_plain_text():
    text = extract_text("resume.txt", SAMPLE_RESUME.encode())
    assert "Lumen Health" in text


def test_extracts_docx_including_tables():
    doc = Document()
    doc.add_paragraph("Priya Raman, software engineer with React and Flask experience across several projects.")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Skills"
    table.rows[0].cells[1].text = "Kubernetes, Terraform"
    buffer = io.BytesIO()
    doc.save(buffer)
    text = extract_text("cv.docx", buffer.getvalue())
    assert "Priya Raman" in text and "Kubernetes" in text


def test_extracts_pdf():
    line = "Backend engineer skilled in Django, PostgreSQL and Docker with five years of experience."
    assert "Django" in extract_text("cv.pdf", _minimal_pdf(line))


@pytest.mark.parametrize(
    ("filename", "data", "code"),
    [
        ("resume.png", b"...", "unsupported_file"),
        ("resume.txt", b"too short", "empty_resume"),
        ("resume.pdf", b"not really a pdf", "invalid_file"),
        ("resume.docx", b"not a zip", "invalid_file"),
    ],
)
def test_rejects_bad_files(filename, data, code):
    with pytest.raises(ResumeError) as err:
        extract_text(filename, data)
    assert err.value.code == code


def test_canonical_names_and_categories():
    assert skills.canonical_name("reactjs") == "React"
    assert skills.canonical_name("  Postgres ") == "PostgreSQL"
    assert skills.canonical_name("AWS (EC2, S3, Lambda)") == "AWS"
    assert skills.canonical_name("Kafka Streams (internal)") == "Kafka Streams (internal)"
    assert skills.category_of("k8s") == "cloud"
    assert skills.category_of("Some Niche Tool") == "other"


def test_scan_text_respects_symbols_and_boundaries():
    found = skills.scan_text("Worked with C++, .NET and Node.js; styled with CSS. Also used React.")
    assert found == ["C++", ".NET", "Node.js", "CSS", "React"]
    # "C" inside "CSS" and "Go" inside "Google" must not count as skills.
    assert "C" not in found
    assert "Go" not in skills.scan_text("Google Cloud and GCP")


def test_merge_skills_dedupes_and_adds_missed_keywords():
    merged = merge_skills(["react.js", "React", "Flask"], "We deployed on Kubernetes.")
    assert [s.name for s in merged] == ["React", "Flask", "Kubernetes"]
    assert merged[2].category == "cloud"


@pytest.mark.parametrize(("years", "level"), [(None, "entry"), (0.5, "entry"), (3, "mid"), (8, "senior")])
def test_seniority(years, level):
    assert seniority_for(years) == level


def test_build_profile_with_llm():
    profile, used_llm = build_profile(SAMPLE_RESUME, FakeLLM({ResumeDraft: resume_draft()}))
    assert used_llm
    assert profile.name == "Priya Raman"
    assert profile.seniority == "entry"
    names = [s.name for s in profile.skills]
    assert names[:2] == ["Python", "React"]
    assert "PostgreSQL" in names  # "Postgres" alias canonicalised


def test_build_profile_falls_back_to_keywords_without_llm():
    profile, used_llm = build_profile(SAMPLE_RESUME, FakeLLM())
    assert not used_llm
    assert {"Python", "React", "Flask", "Docker"} <= {s.name for s in profile.skills}
