"""Turn an uploaded resume file (PDF, DOCX or plain text) into clean text."""

import io
import re
from pathlib import Path

from docx import Document
from pypdf import PdfReader
from pypdf.errors import PdfReadError

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}
# Enough for a multi-page resume while keeping the LLM prompt small.
MAX_CHARS = 12_000


class ResumeError(Exception):
    def __init__(self, message: str, code: str):
        super().__init__(message)
        self.code = code


def _pdf_text(data: bytes) -> str:
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise ResumeError("This PDF is password-protected. Export an unlocked copy and try again.", "invalid_file")
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    except PdfReadError as exc:
        raise ResumeError("This PDF couldn't be read. Try exporting it again or upload a DOCX.", "invalid_file") from exc


def _docx_text(data: bytes) -> str:
    try:
        document = Document(io.BytesIO(data))
    except Exception as exc:  # python-docx raises a variety of zip/xml errors
        raise ResumeError("This DOCX file couldn't be opened. Try saving it again or upload a PDF.", "invalid_file") from exc
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def clean_text(text: str) -> str:
    text = text.replace("\x00", "").replace("•", "-").replace("", "-")
    text = re.sub(r"[ \t ]+", " ", text)
    text = re.sub(r"\n\s*\n\s*\n+", "\n\n", text)
    return text.strip()


def extract_text(filename: str, data: bytes) -> str:
    extension = Path(filename or "").suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise ResumeError("Upload your resume as a PDF, DOCX or TXT file.", "unsupported_file")

    if extension == ".pdf":
        raw = _pdf_text(data)
    elif extension == ".docx":
        raw = _docx_text(data)
    else:
        raw = data.decode("utf-8", errors="replace")

    text = clean_text(raw)
    if len(text) < 80:
        raise ResumeError(
            "Almost no text could be read from this file. If it's a scanned image, upload a text-based PDF or DOCX.",
            "empty_resume",
        )
    return text[:MAX_CHARS]
